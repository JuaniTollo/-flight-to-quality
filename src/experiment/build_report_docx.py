"""Arma un documento .docx con TODAS las figuras del experimento + narrativa científica.

Recorre output/experiment/*.png, las agrupa por pieza con un epígrafe, y produce un
documento legible. Reejecutable: a medida que se generan nuevas figuras (p12_*), se
incorporan automáticamente.

Run:  uv run python -m src.experiment.build_report_docx
"""
from __future__ import annotations

import glob
import os

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

from src.experiment import common as C

OUT = C.OUTDIR / "REPORTE_ciclo_sobreacumulacion.docx"

# epígrafe por figura (científico). Las no listadas se incluyen con su nombre de archivo.
CAPS = {
    "p1_trajectory.png": ("Oscilación característica — trayectoria",
        "Ajuste de los osciladores LV/FN sobre una ventana representativa (par de ciclos)."),
    "p1_phase.png": ("Oscilación característica — retrato de fase", ""),
    "p1_phase_data.png": ("Retrato de fase de los datos", ""),
    "p2_amplitude_boundary.png": ("Frontera de amplitud",
        "Excursión de cada recesión relativa al máximo previo. 2008 redefine la envolvente de posguerra."),
    "p4_vector_field.png": ("Evaluación dinámica — campo de fase",
        "Campo vectorial del oscilador ajustado vs trayectoria observada (gradient matching)."),
    "p5_timefreq.png": ("Contenido tiempo-frecuencia", ""),
    "p5_tvp_resonance.png": ("Oscilador lineal de parámetros variables (VAR rodante)", ""),
    "p7_overaccumulation.png": ("CCF profits↔investment — el patrón empírico",
        "La correlación cruzada oscila (cruza a negativo): co-movimiento contemporáneo dominante "
        "(lag 0) + feedback negativo retardado (lag −4/−5). Robusto sin 2008/2020."),
    "p8_composite_crisis.png": ("Crisis compuesta y buildup de sobreacumulación",
        "Recesiones alineadas en el fondo: la ganancia se quiebra primero y la brecha "
        "inversión−ganancia trepa ~1 año antes del fondo (coincide con Tapia Granados, 2012)."),
    "p8_phase_cycle.png": ("Ciclo estilizado (fase normalizada trough→trough)", ""),
    "p9_crisis_regime.png": ("Régimen de crisis — R² del campo (crisis vs expansión)",
        "Los osciladores no lineales (FN/LV) describen mejor el campo en crisis; el lineal no "
        "distingue régimen. Efecto sugestivo, no significativo (n chico)."),
    "p9_ablation_heatmap.png": ("Ablation del contraste crisis−expansión (288 celdas)",
        "Δ R² por ancla × ancho de ventana; el contraste es positivo salvo con ancla en el pico."),
    "p9_placebo.png": ("Test de placebo / identificabilidad",
        "El ajuste de un oscilador a una ventana de crisis no es especial vs ventanas aleatorias "
        "del mismo largo (confound de amplitud) — el ajuste in-sample no valida el régimen."),
    "p9_pooled_LV.png": ("Dinámica compartida entre crisis — LV", ""),
    "p9_pooled_FN.png": ("Dinámica compartida entre crisis — FN", ""),
    "p11_physical_delay.png": ("Modelo físico de maduración — identificabilidad",
        "El centro del lag de maduración (~1 año) es identificable; el ancho de su distribución "
        "no lo es (mismo ajuste sea distribuido o puntual)."),
}

# orden de presentación (prefijos); el resto va al final por nombre
ORDER = ["p1_", "p2_", "p4_", "p5_", "p7_", "p8_", "p9_", "p11_", "p12_"]

SECTION_INTRO = {
    "p7_": ("1. El patrón empírico (sin modelo)",
        "La correlación cruzada profits↔investment muestra un ciclo endógeno: co-movimiento "
        "contemporáneo dominante y un feedback negativo retardado de la inversión sobre la ganancia."),
    "p8_": ("2. Estructura alrededor de las crisis",
        "Estudio de evento sobre las recesiones: buildup de sobreacumulación ~1 año antes del fondo "
        "y feedback presa-depredador concentrado en la vecindad de las crisis."),
    "p9_": ("3. ¿Es un régimen de crisis? (tests y falsificaciones)",
        "El contraste crisis-vs-expansión es específico de la no-linealidad pero no significativo; "
        "el ajuste in-sample del oscilador a una crisis es, por sí solo, trivial."),
    "p11_": ("4. Modelo físico de maduración (delay distribuido)",
        "Modelo físico diferenciable: la inversión deprime la ganancia tras madurar, con un retardo "
        "distribuido. El centro del lag (~1 año) es identificable; su ancho no."),
    "p12_": ("5. Experimentos complementarios",
        "Calibración por inferencia diferenciable, falsificaciones, identificabilidad y robustez."),
    "p1_": ("Apéndice A. Osciladores: oscilación característica", ""),
    "p2_": ("Apéndice B. Frontera de amplitud", ""),
    "p4_": ("Apéndice C. Evaluación dinámica", ""),
    "p5_": ("Apéndice D. No-estacionariedad", ""),
}


def main():
    pngs = sorted(glob.glob(str(C.OUTDIR / "*.png")))
    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(11)

    doc.add_heading("Ciclo de sobreacumulación profit–investment", level=0)
    p = doc.add_paragraph()
    r = p.add_run("Retardo de maduración distribuido: evidencia empírica y un modelo físico")
    r.italic = True; r.font.size = Pt(12)

    doc.add_heading("Resumen", level=1)
    doc.add_paragraph(
        "El ciclo ganancias–inversión presenta un co-movimiento contemporáneo dominante y un "
        "feedback negativo de la inversión sobre la ganancia con un retardo de maduración "
        "centrado en ~1 año, más intenso cerca de las crisis. Este patrón es capturado por un "
        "modelo físico de cadena de maduración (delay distribuido), no por osciladores de fase "
        "fija (Lotka-Volterra / FitzHugh-Nagumo). El aporte es descriptivo-estructural y de "
        "identificabilidad: el centro del lag se identifica de forma robusta; el ancho de su "
        "distribución y su variación entre crisis no son identificables con los datos disponibles.")

    # agrupar por prefijo en orden
    used = set()
    for pref in ORDER:
        group = [f for f in pngs if os.path.basename(f).startswith(pref)]
        if not group:
            continue
        title, intro = SECTION_INTRO.get(pref, (pref, ""))
        doc.add_heading(title, level=1)
        if intro:
            doc.add_paragraph(intro)
        for f in group:
            used.add(f)
            base = os.path.basename(f)
            cap_title, cap_txt = CAPS.get(base, (base, ""))
            doc.add_heading(cap_title, level=2)
            try:
                doc.add_picture(f, width=Inches(6.3))
                doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
            except Exception as e:
                doc.add_paragraph(f"[no se pudo insertar {base}: {e}]")
            if cap_txt:
                cp = doc.add_paragraph(); cr = cp.add_run(cap_txt)
                cr.italic = True; cr.font.size = Pt(9)

    # cualquier figura no agrupada
    rest = [f for f in pngs if f not in used]
    if rest:
        doc.add_heading("Otras figuras", level=1)
        for f in rest:
            doc.add_heading(os.path.basename(f), level=2)
            try:
                doc.add_picture(f, width=Inches(6.3))
            except Exception:
                pass

    doc.save(str(OUT))
    print(f"✓ {OUT}  ({len(pngs)} figuras)")


if __name__ == "__main__":
    main()
