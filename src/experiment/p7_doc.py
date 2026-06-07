"""Genera el documento metodológico de P7 (CCF profits↔investment) en formato .docx.

Reproducible: usa python-docx + la figura ya generada por p7_overaccumulation.
Run:  uv run python -m src.experiment.p7_doc
"""
from __future__ import annotations

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor, Inches

from src.experiment import common as C

OUT = C.OUTDIR / "p7_metodologia_ccf.docx"
FIG = C.OUTDIR / "p7_overaccumulation.png"


def h(doc, text, level):
    doc.add_heading(text, level=level)


def p(doc, text, italic=False, bold=False):
    par = doc.add_paragraph()
    run = par.add_run(text)
    run.italic = italic
    run.bold = bold
    return par


def code(doc, text):
    par = doc.add_paragraph()
    run = par.add_run(text)
    run.font.name = "Courier New"
    run.font.size = Pt(9)
    par.paragraph_format.left_indent = Inches(0.3)


def main():
    doc = Document()

    # Estilo base
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    # --- Título
    title = doc.add_heading("Ciclo endógeno de sobreacumulación", level=0)
    sub = doc.add_paragraph()
    r = sub.add_run("Metodología de la correlación cruzada profits ↔ investment (Pieza 7)")
    r.italic = True
    r.font.size = Pt(12)
    meta = doc.add_paragraph()
    meta.add_run("Proyecto: ciclo profit–investment · rediseño 2026-06-03").font.size = Pt(9)

    # --- 1. Qué se calcula
    h(doc, "1. Qué es y qué no es", 1)
    p(doc, "Este resultado es estadística descriptiva pura: la función de correlación "
           "cruzada muestral (sample cross-correlation function, CCF) entre dos series "
           "observadas. No interviene ningún modelo —no hay ecuaciones diferenciales, "
           "parámetros ajustados ni osciladores. La estructura cíclica que se reporta está "
           "en los datos, con independencia de cualquier modelo.")
    p(doc, "Esto la distingue de las piezas P1, P3 y P4 del proyecto, que sí ajustan "
           "modelos (Lotka-Volterra, FitzHugh-Nagumo, oscilador lineal) a los datos. P7, "
           "como P2, es evidencia empírica directa.")

    # --- 2. Datos de partida
    h(doc, "2. Datos de partida", 1)
    p(doc, "Dos series trimestrales 1948–2026 (n = 313 trimestres):")
    p(doc, "• PROFITS_YOY — ganancias, variación interanual (YoY)")
    p(doc, "• INVEST_YOY — inversión, variación interanual (YoY)")
    p(doc, "La transformación YoY (hecha en el ETL) convierte niveles con tendencia en "
           "tasas de crecimiento aproximadamente estacionarias. Sin estacionariedad, dos "
           "series con tendencia darían correlación espuria (correlacionarían solo porque "
           "ambas crecen). Esta es la condición que vuelve legítima a la CCF.")

    # --- 3. La definición
    h(doc, "3. Definición de la correlación cruzada", 1)
    p(doc, "Para cada rezago k (en trimestres) se calcula la correlación de Pearson:")
    code(doc, "rho(k) = corr( profits[t] , investment[t + k] )")
    p(doc, "barriendo k de −12 a +12 (±3 años). El signo de k indica quién lidera:")
    p(doc, "• k > 0 → profits comparado con inversión futura → profits lidera")
    p(doc, "• k < 0 → profits comparado con inversión pasada → investment lidera")
    p(doc, "• k = 0 → ambas en el mismo trimestre (co-movimiento contemporáneo)")

    # --- 4. La unidad de k
    h(doc, "4. Por qué k está medido en trimestres", 1)
    p(doc, "La unidad de k la define la frecuencia del dato, no el método. Las series son "
           "trimestrales: cada fila del archivo es un trimestre, identificado por su fecha "
           "de fin de trimestre (DATE). Al pasar a vectores, la posición t en profits[] y "
           "en investment[] corresponde al MISMO trimestre calendario, porque ambos valores "
           "provienen de la misma fila. Por eso desplazar una posición equivale a desplazar "
           "un trimestre, y k = −4 son 4 trimestres = 1 año.")

    # --- 5. k = 0, el ancla
    h(doc, "5. Cómo se ancla k = 0", 1)
    p(doc, "k = 0 no es una elección del método: es el apareamiento natural por fecha "
           "calendario. El archivo trae profits e investment ya fechados por trimestre; "
           "k = 0 los compara en su propia fecha. Ejemplo de las primeras filas:")
    code(doc, "DATE          PROFITS_YOY   INVEST_YOY\n"
              "1948-03-31       6.82         31.65\n"
              "1948-06-30      18.70         45.88\n"
              "1948-09-30      18.13         50.43")
    p(doc, "El par de k = 0 es (profits[1948-Q1], investment[1948-Q1]), ambos de la misma "
           "fecha. El único supuesto metodológico real es que el fechado del dato sea "
           "correcto (que el ETL ubique cada valor en su trimestre verdadero).")

    # --- 6. Ventana de solapamiento
    h(doc, "6. Ventana de solapamiento", 1)
    p(doc, "Para cada k se usan todos los pares solapados, que son n − |k|:")
    code(doc, "k = 0    → 313 pares\n"
              "k = ±4   → 309 pares\n"
              "k = ±12  → 301 pares")
    p(doc, "Los extremos no solapados se descartan. Implicación: a mayor |k|, menos datos y "
           "menor confiabilidad en las puntas. Por eso se acota kmax = 12 (±3 años): "
           "suficiente para ver un ciclo completo sin entrar en lags ruidosos.")

    # --- 7. Preprocesamiento
    h(doc, "7. Estandarización (z-score)", 1)
    p(doc, "Cada serie se estandariza ((x − media) / desvío) antes de correlacionar. Esto "
           "centra en 0 y escala a desvío 1, convirtiendo la covarianza en correlación de "
           "Pearson (adimensional, en [−1, 1]) y comparable entre lags. En P7 el z-score usa "
           "toda la muestra porque el análisis es descriptivo; en P3 (pronóstico) se ajusta "
           "solo con datos de entrenamiento para evitar leakage.")

    # --- 8. Lectura del resultado
    h(doc, "8. Resultado: la CCF oscila (no decae)", 1)
    p(doc, "El criterio que distingue un ciclo endógeno de un mero co-movimiento:")
    p(doc, "• Co-movimiento puro: un pico positivo cerca de k = 0 que decae a cero y se "
           "queda. Las series se mueven juntas, sin retroalimentación.")
    p(doc, "• Ciclo endógeno: la curva cruza a valores negativos y vuelve. Esa inversión de "
           "signo es la firma de un mecanismo que se da vuelta.")
    p(doc, "La CCF observada cruza a negativo en ambos lados → hay ciclo. Valores clave:")
    code(doc, "BOOM  (co-movimiento): lag 0 = +0.67   lag +1 = +0.65\n"
              "GIRO  (sobreacumulación): lag −4 = −0.39  lag −5 = −0.40")
    p(doc, "Interpretación económica (Tapia/Marx):")
    p(doc, "• Boom (lag 0/+1, positivo): las ganancias suben y arrastran la inversión un "
           "trimestre después.")
    p(doc, "• Giro (lag −4/−5, negativo): la inversión de hace ~1 año deprime la ganancia "
           "actual (exceso de acumulación → presión sobre la tasa de ganancia). La demora "
           "de ~1 año refleja el tiempo de maduración de la inversión.")

    # --- Figura
    h(doc, "9. La CCF graficada", 1)
    if FIG.exists():
        doc.add_picture(str(FIG), width=Inches(6.2))
        cap = doc.add_paragraph()
        cr = cap.add_run("Figura 1. CCF profits↔investment. Barras: serie completa "
                         "1948–2026. Línea verde: sin 2008/2020 (robustez). Banda roja: "
                         "zona de sobreacumulación (lags negativos, correlación negativa).")
        cr.italic = True
        cr.font.size = Pt(9)
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # --- 10. Robustez
    h(doc, "10. Robustez", 1)
    p(doc, "La CCF se recomputa excluyendo 2008 y 2020. Si la oscilación fuera artefacto de "
           "las crisis, desaparecería; en cambio el patrón se mantiene o se profundiza:")
    code(doc, "                con crisis    sin 2008/2020\n"
              "lag 0             +0.67          +0.74\n"
              "lag −4            −0.39          −0.40\n"
              "lag +6            −0.22          −0.30")
    p(doc, "El ciclo de sobreacumulación es por tanto estructura de la economía normal, no "
           "efecto de outliers.")

    # --- 11. Matiz
    h(doc, "11. Matiz importante", 1)
    p(doc, "Aunque la CCF tiene la forma de un ciclo presa-depredador, el componente "
           "dominante es el co-movimiento contemporáneo: el pico de +0.67 es bastante mayor "
           "que los valles de −0.40. La fase efectiva del ciclo es ~10° (casi co-movimiento, "
           "profits liderando ~1 trimestre), no los 90° que imponen los osciladores "
           "Lotka-Volterra / FitzHugh-Nagumo. Por eso la firma cíclica existe (sirve para "
           "describir) pero es demasiado suave para que un oscilador no lineal la explote "
           "mejor que un modelo lineal en pronóstico (resultado de P3/P4).")

    # --- Síntesis
    h(doc, "12. Síntesis", 1)
    p(doc, "P7 prueba, solo con datos y sin modelo, que existe un ciclo endógeno de "
           "sobreacumulación entre ganancias e inversión: real, robusto y con la estructura "
           "temporal esperada por la teoría. Es el hallazgo positivo central del proyecto y "
           "el cimiento empírico sobre el cual los modelos no lineales (P1/P3/P4) luego "
           "fracasan en mejorar al baseline lineal.")

    doc.save(str(OUT))
    print(f"✓ documento generado: {OUT}")


if __name__ == "__main__":
    main()
